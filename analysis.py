#!/usr/bin/env python3
"""
Reproducible analysis: Women in STEM/IT in Kyrgyzstan (survey, n≈61)
--------------------------------------------------------------------
Usage:   python analysis.py path/to/raw_export.xlsx [output_dir]
Creates: data/survey_clean_anonymized.xlsx   (cleaned data + codebook, contacts removed)
         data/specialty_coding.csv           (STEM/non-STEM coding of specialties; manual_code overrides auto_code)
         results/results_tables.xlsx         (all tables)
         figures/*.png                       (all charts)
         results/run_log.txt                 (checks + software versions)
The raw file is NOT part of the package because it contains contact details.
"""
import sys, re, platform
from pathlib import Path
import numpy as np
import pandas as pd
from scipy import stats
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

RAW = Path(sys.argv[1]) if len(sys.argv) > 1 else None
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(__file__).parent
if RAW is None or not RAW.exists():
    sys.exit("Give the path to the raw Google-Forms export, e.g.  python analysis.py raw.xlsx")
for d in ("data", "results", "figures"): (OUT/d).mkdir(parents=True, exist_ok=True)
LOG = []
def log(*a):
    s = " ".join(str(x) for x in a); print(s); LOG.append(s)

# ------------------------------------------------------------------ 1. LOAD
raw = pd.read_excel(RAW)
raw.columns = [str(c).strip() for c in raw.columns]
assert raw.shape[1] == 44, f"Expected 44 columns, found {raw.shape[1]}"
assert "Контакты" in raw.columns[43], "Column 44 should be the contacts column"
log(f"Loaded {raw.shape[0]} responses x {raw.shape[1]} columns")

# (position, short name, type, description / coding)   -- position = column index in the raw file
SPEC = [
 (1,"age_group","cat","Age group"),(2,"childhood_place","cat","Where childhood was spent"),
 (3,"marital_status","cat","Marital status"),(4,"has_children","bin","Has children (1=yes)"),
 (5,"mother_edu","cat","Mother's education"),(6,"father_edu","cat","Father's education"),
 (7,"siblings","num","Number of siblings (free text parsed to number; 'Нас 5' read as 4 siblings)"),
 (8,"birth_order","cat","Birth order in family"),
 (9,"family_finance","lik","Family financial situation in childhood, self-rated 1-5"),
 (10,"family_religiosity","lik","Family religiosity, self-rated 1-5 (higher = more religious)"),
 (11,"highest_edu","cat","Highest education"),(12,"specialty_text","text","Specialty (free text)"),
 (13,"stem_class","bin","Attended specialised STEM class/programme (1=yes)"),
 (14,"olympiad","bin","Ever took part in a STEM olympiad/competition (1=yes)"),
 (15,"olympiad_level","cat","Highest olympiad level (only if olympiad=yes)"),
 (16,"why_not","cat","Reason for not taking part (asked if 'No' to STEM class or olympiad)"),
 (17,"role_model","bin","Had a female STEM/IT role model in childhood (1=yes)"),
 (18,"role_model_who","text","Who the role model was (free text)"),
 (19,"parents_support","lik","Parents actively supported my STEM/IT aspirations (1=strongly disagree … 5=strongly agree)"),
 (20,"teacher_support","lik","Teachers actively supported my STEM aspirations (same scale)"),
 (21,"teacher_suggested","lik","Teachers specifically suggested olympiads/competitions (same scale)"),
 (22,"family_suitable","lik","Family/environment saw STEM/IT as 'suitable' for a woman (same scale)"),
 (23,"belonging","lik","I felt I belonged in classes/spaces with more boys (same scale; POSITIVELY worded)"),
 (24,"self_doubt","lik","I doubted my STEM ability more than male classmates (same scale)"),
 (25,"few_women_doubt","lik","Few women around made me doubt continuing (same scale)"),
 (26,"moment_text","text","A specific moment (free text)"),
 (27,"employment","cat","Current employment status"),(28,"job_text","text","Position / sector / tenure (free text)"),
 (29,"interview_family_q","bin","Asked at interview about marriage/pregnancy/children (1=yes; 'n/a' → missing)"),
 (30,"maternity_leave","bin","Took maternity leave (1=yes; 'n/a' → missing)"),
 (31,"maternity_career_impact","bin","This noticeably affected career growth (1=yes; 'n/a' → missing)"),
 (32,"colleague_behaviour","multi","Experiences with male colleagues (multi-select; see beh_* columns)"),
 (33,"pay_fair","lik","I am paid fairly compared with male peers (same scale; workers only)"),
 (34,"family_status_affected","lik","Family status/children affected hiring/promotion (same scale)"),
 (35,"prove_competence","lik","Had to prove competence more than male colleagues (same scale)"),
 (36,"anti_disc_policy","lik","Workplace has clear, enforced anti-discrimination policies (same scale)"),
 (37,"salary_vs_men","cat","Salary relative to male colleagues ('n/a' → missing)"),
 (38,"top3_measures","multi","Top-3 most effective measures (multi-select; see meas_* columns)"),
 (39,"mentor","bin","Had a formal/informal mentor (1=yes)"),
 (40,"mentor_would_change","lik","A mentor would have changed my path (same scale)"),
 (41,"women_leaders_inspire","lik","More women in leadership would draw more girls to STEM (same scale)"),
 (42,"one_change","cat","ONE change that would matter most to me"),
]
df = pd.DataFrame({"resp_id": [f"R{i+1:02d}" for i in range(len(raw))]})
for pos, name, typ, desc in SPEC:
    s = raw.iloc[:, pos].copy()
    if typ in ("cat","text","multi"):
        s = s.astype("object").where(s.notna(), np.nan)
        s = s.map(lambda v: v.strip() if isinstance(v, str) else v)
    df[name] = s

NA_TOK = {"не применимо", "неприменимо"}
def na_out(s): return s.where(~s.astype(str).str.strip().str.lower().isin(NA_TOK), np.nan)
def yn(s): return na_out(s).map({"Да": 1, "Нет": 0}).astype(float)

for c in ["has_children","stem_class","olympiad","role_model","mentor",
          "interview_family_q","maternity_leave","maternity_career_impact"]:
    df[c] = yn(df[c])
for c in ["salary_vs_men"]: df[c] = na_out(df[c])
for c,_ in [(x[1],0) for x in SPEC if x[2]=="lik"]: df[c] = pd.to_numeric(df[c], errors="coerce")
df["employment"] = df["employment"].map(lambda v: v[0].upper()+v[1:] if isinstance(v,str) and v else v)

def parse_sib(v):
    if pd.isna(v): return np.nan
    if isinstance(v,(int,float)): return float(v)
    s = str(v).strip().lower(); n = re.findall(r"\d+", s)
    if not n: return np.nan
    return float(n[0])-1 if s.startswith("нас") else float(sum(int(x) for x in n))
sib_raw = df["siblings"].copy(); df["siblings"] = sib_raw.map(parse_sib)
sib_manual = [(r, v) for r, v in zip(df.resp_id, sib_raw) if isinstance(v, str)]
log(f"Siblings: {len(sib_manual)} free-text entries parsed to numbers:", sib_manual)

# derived groupings
df["age_lt25"] = df["age_group"].isin(["До 18","18–24"]).astype(int)
df["adult"] = (df["age_group"] != "До 18").astype(int)
df["village"] = (df["childhood_place"] == "Село").astype(float)
df["bishkek"] = (df["childhood_place"] == "Столица (Бишкек)").astype(float)
df["emp_group"] = df["employment"].map(lambda v: "IT/STEM workers" if v=="Работаю в IT/STEM"
                    else "Non-STEM workers" if v=="Работаю вне IT/STEM" else "Other (students, job-seekers, not working)")

# multi-select: colleague behaviour (substring matching – option labels contain commas!)
BEH = {"beh_interrupted":["перебивал","не давали договорить"],"beh_ideas_credited":["приписывали"],
       "beh_excluded_informal":["неформального"],"beh_skills_underestimated":["недооценивали"],
       "beh_harassment_comments":["неуместные комментарии"],"beh_none":["ничего из перечисленного"]}
beh_app = df["colleague_behaviour"].notna() & ~df["colleague_behaviour"].astype(str).str.strip().str.lower().isin(NA_TOK)
for k, subs in BEH.items():
    df[k] = [np.nan if not a else float(any(t in str(v).lower() for t in subs))
             for v, a in zip(df["colleague_behaviour"], beh_app)]
# multi-select: top-3 measures
MEAS = {"meas_equal_pay":"равная","meas_flexible_childcare":"гибкий график","meas_leadership_confidence":"лидерства",
        "meas_mentoring":"менторства","meas_teacher_training":"учителей","meas_workplace_culture":"культуре и союзничеству",
        "meas_girls_only_stem":"только для девочек"}
for k, sub in MEAS.items():
    df[k] = df["top3_measures"].astype(str).str.lower().map(lambda v: float(sub in v))
df["n_measures_chosen"] = df[list(MEAS)].sum(axis=1)
log("Respondents choosing more than 3 measures:", int((df.n_measures_chosen > 3).sum()), "(kept as given)")

# specialty → STEM coding (auto by keywords, editable via CSV)
KW = ["engineer","инжен","программ","data science","математ","информац","иит","tech","компютер",
      "энергет","software","mechanical","клеточная","эколог"]
NO_SPEC = {"школа","общее среднее образование","нету пока"}
REVIEW = ["mpa in tech","nursing","компютер курс","училась на первом курсе","эколог","дизайн архитектурной"]
csv_path = OUT/"data"/"specialty_coding.csv"
auto = []
for rid, t, ad in zip(df.resp_id, df.specialty_text.fillna(""), df.adult):
    tl = t.strip().lower()
    code = np.nan if (not ad or tl in NO_SPEC or tl == "") else float(any(k in tl for k in KW))
    auto.append((rid, t.strip(), code, int(any(k in tl for k in REVIEW))))
coding = pd.DataFrame(auto, columns=["resp_id","specialty_text","auto_code","needs_review"])
if csv_path.exists():
    old = pd.read_csv(csv_path)
    coding = coding.merge(old[["resp_id","manual_code"]], on="resp_id", how="left")
else:
    coding["manual_code"] = np.nan
coding.to_csv(csv_path, index=False)
coding["final"] = coding["manual_code"].where(coding["manual_code"].notna(), coding["auto_code"])
df["stem_field"] = coding["final"].values
log("STEM field (adults with a coded specialty):", int(df.stem_field.notna().sum()),
    "| STEM:", int((df.stem_field==1).sum()), "| flagged needs_review:", int(coding.needs_review.sum()))

# ------------------------------------------------------------------ 2. SAVE CLEAN DATA + CODEBOOK
free_cols = ["specialty_text","role_model_who","moment_text","job_text"]
codebook = pd.DataFrame([(n, raw.columns[p][:120], t, d) for p,n,t,d in SPEC],
                        columns=["variable","original_question (Russian, truncated)","type","description / coding"])
extra = [("resp_id","(generated)","id","Anonymous respondent id, order of rows in export"),
 ("adult","(derived)","bin","1 = age 18+"),("age_lt25","(derived)","bin","1 = under 18 or 18–24"),
 ("village","(derived)","bin","1 = childhood in a village"),("bishkek","(derived)","bin","1 = childhood in Bishkek"),
 ("emp_group","(derived)","cat","IT/STEM workers / Non-STEM workers / Other"),
 ("stem_field","(derived)","bin","1 = adult whose specialty is STEM/IT (see data/specialty_coding.csv); missing for minors and school-only"),
 ("beh_*","(derived)","bin","1/0 per colleague-behaviour option among those it applies to (n/a → missing)"),
 ("meas_*","(derived)","bin","1/0 per measure selected in top-3 list"),("n_measures_chosen","(derived)","num","How many measures were ticked (>3 means the limit was exceeded)")]
codebook = pd.concat([codebook, pd.DataFrame(extra, columns=codebook.columns)], ignore_index=True)
with pd.ExcelWriter(OUT/"data"/"survey_clean_anonymized.xlsx") as w:
    df.drop(columns=free_cols+["colleague_behaviour","top3_measures"]).to_excel(w, sheet_name="data", index=False)
    codebook.to_excel(w, sheet_name="codebook", index=False)
    df[["resp_id"]+free_cols].to_excel(w, sheet_name="free_text_REVIEW_NAMES", index=False)

# ------------------------------------------------------------------ 3. HELPERS
def pct(x, n): return round(100*x/n, 1) if n else np.nan
def likert_row(s, label):
    s = s.dropna(); n = len(s)
    r = {"item": label, "n": n}
    for k in range(1,6): r[f"%{k}"] = pct((s==k).sum(), n)
    r.update({"median": s.median(), "%agree (4-5)": pct((s>=4).sum(), n), "%disagree (1-2)": pct((s<=2).sum(), n)})
    return r
def mw(x1, x0):
    x1, x0 = x1.dropna(), x0.dropna()
    if len(x1) < 3 or len(x0) < 3: return None
    u, p = stats.mannwhitneyu(x1, x0, alternative="two-sided")
    return dict(n_group1=len(x1), n_group0=len(x0), summary=f"median {x1.median():g} vs {x0.median():g}",
                effect=round(2*u/(len(x1)*len(x0))-1, 3), effect_type="rank-biserial r", p=p)
def fisher(pred, out):
    m = pred.notna() & out.notna()
    t = pd.crosstab(pred[m], out[m]).reindex(index=[0,1], columns=[0,1], fill_value=0)
    a, b, c, d = t.loc[1,1], t.loc[1,0], t.loc[0,1], t.loc[0,0]
    p = stats.fisher_exact(t.values)[1]
    odds = ((a+.5)*(d+.5))/((b+.5)*(c+.5))
    n1, n0 = a+b, c+d
    return dict(n_group1=int(n1), n_group0=int(n0),
                summary=f"{pct(a,n1)}% ({a}/{n1}) vs {pct(c,n0)}% ({c}/{n0})",
                effect=round(odds,2), effect_type="odds ratio (Haldane-corrected)", p=p)
def holm(ps):
    ps = np.array(ps, float); o = np.argsort(ps); m = len(ps); adj = np.empty(m); run = 0
    for r, i in enumerate(o):
        run = max(run, (m-r)*ps[i]); adj[i] = min(1, run)
    return adj
LIK_PRED = {"family_finance":"Family finances (childhood)","family_religiosity":"Family religiosity",
    "parents_support":"Parents supported STEM","teacher_support":"Teachers supported STEM",
    "teacher_suggested":"Teachers suggested olympiads","family_suitable":"Family saw STEM as suitable for women",
    "belonging":"Felt I belonged among boys","self_doubt":"Doubted ability more than boys","few_women_doubt":"Few women made me doubt"}
BIN_PRED = {"role_model":"Had female role model","mentor":"Had a mentor","village":"Village childhood","bishkek":"Bishkek childhood"}
OUTS = {"stem_class":"Attended STEM class","olympiad":"Took part in STEM olympiad","stem_field":"STEM field of study (adults)"}
def run_rq1(data, outs):
    rows = []
    for o, ol in outs.items():
        for k, kl in LIK_PRED.items():
            r = mw(data.loc[data[o]==1, k], data.loc[data[o]==0, k])
            if r: rows.append({"outcome": ol, "factor": kl, "test": "Mann-Whitney (factor score: outcome yes vs no)", **r})
        for k, kl in BIN_PRED.items():
            r = fisher(data[k], data[o])
            rows.append({"outcome": ol, "factor": kl, "test": "Fisher exact (% with outcome: factor yes vs no)", **r})
    t = pd.DataFrame(rows); t["p_holm"] = holm(t.p)
    t["p"] = t.p.round(4); t["p_holm"] = t.p_holm.round(4)
    t["flag"] = np.where(t.p_holm < .05, "survives Holm", np.where(t.p < .05, "p<.05 only (exploratory)", ""))
    return t

# ------------------------------------------------------------------ 4. TABLES
T = {}
# sample description
rows = []
for v, lab in [("age_group","Age group"),("childhood_place","Childhood place"),("marital_status","Marital status"),
               ("has_children","Has children"),("highest_edu","Highest education"),("employment","Employment")]:
    vc = df[v].value_counts(dropna=False)
    for k, n in vc.items(): rows.append({"variable": lab, "category": k, "n": n, "%": pct(n, len(df))})
T["1_sample"] = pd.DataFrame(rows)
# Likert overview
LOV = [("parents_support","Parents supported STEM"),("teacher_support","Teachers supported STEM"),
       ("teacher_suggested","Teachers suggested olympiads"),("family_suitable","Family saw STEM as suitable"),
       ("belonging","Felt I belonged among boys"),("self_doubt","Doubted ability more than boys"),
       ("few_women_doubt","Few women made me doubt"),("family_religiosity","Family religiosity"),
       ("family_finance","Family finances"),("pay_fair","[work] Paid fairly vs men"),
       ("family_status_affected","[work] Family status affected hiring/promotion"),
       ("prove_competence","[work] Had to prove competence more"),("anti_disc_policy","[work] Enforced anti-discrimination policy"),
       ("mentor_would_change","Mentor would have changed my path"),("women_leaders_inspire","Women leaders would draw girls")]
T["2_likert_overview"] = pd.DataFrame([likert_row(df[k], l) for k, l in LOV])
# RQ1
T["3_rq1_outcomes"] = pd.DataFrame([{"outcome": OUTS[o], "yes_n": int((df[o]==1).sum()), "no_n": int((df[o]==0).sum()),
    "%yes": pct((df[o]==1).sum(), df[o].notna().sum())} for o in OUTS])
r_why = df.loc[df.why_not.notna() & ~df.why_not.astype(str).str.startswith("Неприменимо"), "why_not"].value_counts()
WHY = {"Не было интереса в тот момент":"No interest at the time","Не знала, что это существует":"Did not know it existed",
       "Не было доступа в моём регионе":"No access in my region","Никто не приглашал и не номинировал":"Nobody invited/nominated me",
       "Учителя не поддерживали":"Teachers did not support","Семья не поддерживала":"Family did not support"}
T["4_rq1_reasons_not_taking_part"] = pd.DataFrame({"reason": [WHY.get(k,k) for k in r_why.index], "n": r_why.values,
    "%": [pct(n, r_why.sum()) for n in r_why.values]})
T["5_rq1_factor_tests"] = run_rq1(df, OUTS)
T["6_rq1_robust_adults_only"] = run_rq1(df[df.adult==1], OUTS)
lik_cols = list(LIK_PRED)
rho = df[lik_cols].corr(method="spearman").round(2); rho.index = rho.columns = [LIK_PRED[c] for c in lik_cols]
pv = pd.DataFrame(index=lik_cols, columns=lik_cols, dtype=float)
for a in lik_cols:
    for b in lik_cols: pv.loc[a,b] = stats.spearmanr(df[a], df[b])[1] if a != b else np.nan
pv = pv.round(4); pv.index = pv.columns = rho.columns
T["7_rq1_spearman_rho"] = rho.reset_index().rename(columns={"index":""}); T["8_rq1_spearman_p"] = pv.reset_index().rename(columns={"index":""})
def alpha(d):
    d = d.dropna(); k = d.shape[1]
    return round(k/(k-1)*(1 - d.var(ddof=1).sum()/d.sum(axis=1).var(ddof=1)), 3), len(d)
a1 = alpha(df[["parents_support","teacher_support","teacher_suggested","family_suitable"]])
a2 = alpha(pd.concat([6-df["belonging"], df[["self_doubt","few_women_doubt"]]], axis=1))
T["9_cronbach_alpha"] = pd.DataFrame([
    {"scale": "Support (parents, teachers, teacher suggestions, family suitability)", "items": 4, "alpha": a1[0], "n_complete": a1[1]},
    {"scale": "Insecurity (belonging REVERSED, self-doubt, few-women doubt)", "items": 3, "alpha": a2[0], "n_complete": a2[1]}])
T["10_age_cohorts"] = df.groupby("age_group")[["stem_class","olympiad"]].agg(["sum","count","mean"]).round(2).reset_index()
T["10_age_cohorts"].columns = ["age_group","stem_class_n","stem_class_N","stem_class_share","olymp_n","olymp_N","olymp_share"]
# RQ2
groups = {"IT/STEM workers": df.emp_group=="IT/STEM workers", "Non-STEM workers": df.emp_group=="Non-STEM workers",
          "Other who answered": (df.emp_group.str.startswith("Other")) & df.pay_fair.notna(), "All who answered": df.pay_fair.notna()}
rows = []
for g, m in groups.items():
    for v, lab in [("interview_family_q","Asked about marriage/children at interview"),("maternity_leave","Took maternity leave"),
                   ("maternity_career_impact","Leave/family affected career growth")]:
        s = df.loc[m, v].dropna(); rows.append({"group": g, "item": lab, "applicable_n": len(s), "yes_n": int(s.sum()), "%yes": pct(s.sum(), len(s))})
    s = df.loc[m, "salary_vs_men"].dropna()
    for k, n in s.value_counts().items(): rows.append({"group": g, "item": f"Salary vs men: {k}", "applicable_n": len(s), "yes_n": n, "%yes": pct(n, len(s))})
T["11_rq2_yes_no_items"] = pd.DataFrame(rows)
T["12_rq2_likert_by_group"] = pd.DataFrame([{"group": g, **likert_row(df.loc[m, k], l)} for g, m in groups.items()
    for k, l in [("pay_fair","Paid fairly vs men"),("family_status_affected","Family status affected hiring/promotion"),
                 ("prove_competence","Had to prove competence more"),("anti_disc_policy","Enforced anti-discrimination policy")]])
BEHL = {"beh_skills_underestimated":"Technical skills underestimated","beh_harassment_comments":"Inappropriate comments/harassment",
        "beh_interrupted":"Interrupted / not allowed to finish","beh_ideas_credited":"Ideas credited to a man",
        "beh_excluded_informal":"Excluded from informal networking","beh_none":"None of these"}
rows = []
for g in ["IT/STEM workers","Non-STEM workers","All who answered"]:
    m = groups[g] if g != "All who answered" else df.colleague_behaviour.notna()
    for k, l in BEHL.items():
        s = df.loc[m, k].dropna(); rows.append({"group": g, "experience": l, "applicable_n": len(s), "n": int(s.sum()), "%": pct(s.sum(), len(s))})
T["13_rq2_colleague_behaviour"] = pd.DataFrame(rows)
# RQ3
MEASL = {"meas_leadership_confidence":"Leadership & confidence programmes","meas_equal_pay":"Equal / transparent pay",
 "meas_flexible_childcare":"Flexible schedule / childcare support","meas_teacher_training":"Teacher training on gender bias",
 "meas_mentoring":"Mentoring programmes","meas_workplace_culture":"Workplace culture / allyship training","meas_girls_only_stem":"Girls-only STEM programmes"}
T["14_rq3_top3_measures"] = pd.DataFrame([{"measure": l, "n": int(df[k].sum()), f"% of all (n={len(df)})": pct(df[k].sum(), len(df)),
    "% under 25": pct(df.loc[df.age_lt25==1, k].sum(), (df.age_lt25==1).sum()), "% 25+": pct(df.loc[df.age_lt25==0, k].sum(), (df.age_lt25==0).sum())}
    for k, l in MEASL.items()]).sort_values("n", ascending=False)
CH = {"Больше уверенности в себе с раннего возраста":"More self-confidence from early age","Доступная менторская программа":"Accessible mentoring programme",
 "Более безопасная и уважительная рабочая среда":"Safer, more respectful workplace","Больше поддержки от учителей/наставников":"More support from teachers",
 "Больше поддержки от родителей/семьи":"More support from parents/family","Наличие ролевых моделей — успешных женщин в STEM/IT":"Female role models",
 "Доступ к курсам/олимпиадам для девочек":"Access to courses/olympiads for girls","Гибкий график / поддержка с детьми":"Flexible schedule / childcare",
 "Равная оплата труда":"Equal pay","Информированность":"Awareness/information"}
vc = df.one_change.value_counts()
T["15_rq3_one_change"] = pd.DataFrame({"change": [CH.get(k,k) for k in vc.index], "n": vc.values, "%": [pct(n, len(df)) for n in vc.values]})
# data quality
dq = pd.DataFrame({"variable": df.columns, "missing_n": df.isna().sum().values})
T["16_data_quality"] = dq[dq.missing_n > 0]

# ------------------------------------------------------------------ 5. WRITE RESULTS WORKBOOK
README = pd.DataFrame({"Notes": [
 "All numbers in this workbook are produced by analysis.py from the raw export (statistical tests cannot be expressed as spreadsheet formulas; re-run the script to reproduce).",
 "Sample: non-random (self-selected) survey; results describe respondents, not all women in Kyrgyzstan. Associations, not causes.",
 "Likert items: 1 = strongly disagree … 5 = strongly agree, treated as ordinal (medians, % agree 4-5, % disagree 1-2).",
 "Mann-Whitney: compares factor scores between people with vs without the outcome; effect = rank-biserial r (−1…+1, 0 = none).",
 "Fisher exact: compares % with the outcome between factor yes/no; effect = odds ratio with 0.5 added to each cell.",
 "p_holm = Holm-adjusted p across all tests in that table. 'survives Holm' is the conservative flag; 'p<.05 only' is exploratory.",
 "STEM field of study is coded from free-text answers by keywords; ambiguous answers are flagged for manual review. Adults with a stated field only.",
 "'n/a' answers are treated as missing; every percentage states its denominator (n).",
 "Sheet 5 has 39 tests, sheet 6 repeats them for adults only as a robustness check."]})
xl = OUT/"results"/"results_tables.xlsx"
with pd.ExcelWriter(xl) as w:
    README.to_excel(w, sheet_name="README", index=False)
    for k, t in T.items(): t.to_excel(w, sheet_name=k[:31], index=False)
wb = load_workbook(xl)
for ws in wb:
    for row in ws.iter_rows():
        for c in row: c.font = Font(name="Arial", size=10)
    for c in ws[1]: c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
    ws.freeze_panes = "A2"
    for i, col in enumerate(ws.columns, 1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width+2), 90 if ws.title=="README" else 60)
    if ws.title == "README":
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
wb.save(xl)

# ------------------------------------------------------------------ 6. FIGURES
plt.rcParams.update({"font.size": 9, "axes.spines.top": False, "axes.spines.right": False})
C_DIS, C_NEU, C_AGR = "#D55E00", "#BDBDBD", "#0072B2"
def save(fig, name): fig.tight_layout(); fig.savefig(OUT/"figures"/name, dpi=200, bbox_inches="tight"); plt.close(fig)
# Fig 1 diverging Likert
items = [(k, l) for k, l in LOV if k not in ("family_religiosity","family_finance")]
fig, ax = plt.subplots(figsize=(8.5, 5.2))
for i, (k, l) in enumerate(reversed(items)):
    s = df[k].dropna(); n = len(s); d, ne, a = 100*(s<=2).mean(), 100*(s==3).mean(), 100*(s>=4).mean()
    ax.barh(i, -d, color=C_DIS); ax.barh(i, -ne/2, left=-d, color=C_NEU); ax.barh(i, ne/2, color=C_NEU); ax.barh(i, a, left=ne/2, color=C_AGR)
    ax.text(-d-ne/2-1, i, f"{d:.0f}%", ha="right", va="center", fontsize=8); ax.text(a+ne/2+1, i, f"{a:.0f}%", ha="left", va="center", fontsize=8)
ax.set_yticks(range(len(items))); ax.set_yticklabels([f"{l} (n={df[k].notna().sum()})" for k, l in reversed(items)])
ax.axvline(0, color="k", lw=.6); ax.set_xlim(-90, 95); ax.set_xlabel("% of respondents  (orange = disagree 1–2, grey = neutral 3, blue = agree 4–5)")
ax.set_title("Fig 1. Agreement with statements about school, family and work"); save(fig, "fig1_likert_overview.png")
# Fig 2 outcomes by level of teacher/parent/belonging
lv = lambda s: pd.cut(s, [0,2,3,5], labels=["Low (1–2)","Mid (3)","High (4–5)"])
fig, axs = plt.subplots(1, 3, figsize=(10, 3.6), sharey=True)
for ax, (k, l) in zip(axs, [("teacher_support","Teacher support"),("parents_support","Parental support"),("belonging","Sense of belonging")]):
    g = df.assign(lvl=lv(df[k]))
    for j, (o, ol, col) in enumerate([("olympiad","Olympiad","#0072B2"),("stem_field","STEM field (adults)","#E69F00")]):
        pr = [100*g.loc[g.lvl==L, o].mean() for L in g.lvl.cat.categories]; ns = [g.loc[(g.lvl==L)&g[o].notna(), o].shape[0] for L in g.lvl.cat.categories]
        xs = np.arange(3)+(j-.5)*.38; ax.bar(xs, pr, .38, color=col, label=ol if k=="teacher_support" else None)
        for x, p_, n in zip(xs, pr, ns): ax.text(x, p_+1.5, f"n={n}", ha="center", fontsize=7)
    ax.set_xticks(range(3)); ax.set_xticklabels(["Low\n(1–2)","Mid\n(3)","High\n(4–5)"]); ax.set_title(l)
axs[0].set_ylabel("% who took part / study STEM"); axs[0].legend(frameon=False, fontsize=8, loc="upper left")
fig.suptitle("Fig 2. Olympiad participation and STEM study by level of support and belonging", y=1.04); save(fig, "fig2_outcomes_by_support.png")
# Fig 3 reasons
t = T["4_rq1_reasons_not_taking_part"].iloc[::-1]
fig, ax = plt.subplots(figsize=(7, 3)); ax.barh(t.reason, t["%"], color="#0072B2")
for i, (p_, n) in enumerate(zip(t["%"], t.n)): ax.text(p_+.6, i, f"{p_:.0f}% (n={n})", va="center", fontsize=8)
ax.set_xlim(0, t["%"].max()*1.3); ax.set_xlabel("% of those giving a reason"); ax.set_title(f"Fig 3. Why not a STEM class / olympiad? (n={int(t.n.sum())})"); save(fig, "fig3_reasons_not_taking_part.png")
# Fig 4 mentor & role model
fig, axs = plt.subplots(1, 2, figsize=(9, 3.6), sharey=True)
for ax, (k, l) in zip(axs, [("mentor","Had a mentor"),("role_model","Had a female role model")]):
    for j, (o, ol, col) in enumerate([("stem_class","STEM class","#009E73"),("olympiad","Olympiad","#0072B2"),("stem_field","STEM field (adults)","#E69F00")]):
        for xi, val in enumerate([0, 1]):
            m = (df[k]==val) & df[o].notna(); p_ = 100*df.loc[m, o].mean(); x = xi+(j-1)*.27
            ax.bar(x, p_, .27, color=col, label=ol if (k=="mentor" and xi==0) else None); ax.text(x, p_+1.5, f"{m.sum()}", ha="center", fontsize=7)
    ax.set_xticks([0,1]); ax.set_xticklabels(["No","Yes"]); ax.set_title(l)
axs[0].set_ylabel("% with outcome (numbers = group n)"); axs[0].legend(frameon=False, fontsize=8, loc="upper left")
fig.suptitle(f"Fig 4. Mentors and role models (small groups: {int((df.mentor==1).sum())} mentors, {int((df.role_model==1).sum())} role models)", y=1.04); save(fig, "fig4_mentor_role_model.png")
# Fig 5 workplace behaviour
t = T["13_rq2_colleague_behaviour"]; t = t[t.group.isin(["IT/STEM workers","Non-STEM workers"]) & (t.experience!="None of these")]
order = list(BEHL.values())[:-1][::-1]
fig, ax = plt.subplots(figsize=(8, 3.8))
for j, (g, col) in enumerate([("IT/STEM workers","#0072B2"),("Non-STEM workers","#E69F00")]):
    tt = t[t.group==g].set_index("experience").loc[order]; y = np.arange(len(order))+(j-.5)*.38
    ax.barh(y, tt["%"], .38, color=col, label=f"{g} (n={int(tt.applicable_n.iloc[0])})")
    for yy, p_, n in zip(y, tt["%"], tt.n): ax.text(p_+1, yy, f"{p_:.0f}% ({int(n)})", va="center", fontsize=7)
ax.set_yticks(range(len(order))); ax.set_yticklabels(order); ax.set_xlim(0, t["%"].max()*1.3); ax.legend(frameon=False, fontsize=8)
ax.set_xlabel("% of those with work experience"); ax.set_title("Fig 5. Experiences with male colleagues\n(very small groups – descriptive only)"); save(fig, "fig5_workplace_behaviour.png")
# Fig 6/7 RQ3
t = T["14_rq3_top3_measures"].iloc[::-1]; col = [c for c in t.columns if c.startswith("% of all")][0]
fig, ax = plt.subplots(figsize=(7.5, 3.6)); ax.barh(t.measure, t[col], color="#0072B2")
for i, (p_, n) in enumerate(zip(t[col], t.n)): ax.text(p_+.8, i, f"{p_:.0f}% ({n})", va="center", fontsize=8)
ax.set_xlim(0, t[col].max()*1.3); ax.set_xlabel("% of respondents"); ax.set_title(f"Fig 6. Measures chosen among the top 3 (n={len(df)}; perceptions, not evidence of effect)"); save(fig, "fig6_rq3_top3_measures.png")
t = T["15_rq3_one_change"].iloc[::-1]
fig, ax = plt.subplots(figsize=(7.5, 3.8)); ax.barh(t.change, t["%"], color="#009E73")
for i, (p_, n) in enumerate(zip(t["%"], t.n)): ax.text(p_+.6, i, f"{p_:.0f}% ({n})", va="center", fontsize=8)
ax.set_xlim(0, t["%"].max()*1.3); ax.set_xlabel("% of respondents"); ax.set_title(f"Fig 7. The ONE change that would matter most to me (n={len(df)})"); save(fig, "fig7_rq3_one_change.png")

# ------------------------------------------------------------------ 7. LOG
log("Python", platform.python_version(), "| pandas", pd.__version__, "| scipy", __import__("scipy").__version__)
log("Tests run in sheet 5:", len(T["5_rq1_factor_tests"]), "| significant after Holm:", int((T["5_rq1_factor_tests"].p_holm < .05).sum()),
    "| p<.05 unadjusted:", int((T["5_rq1_factor_tests"].p < .05).sum()))
(OUT/"results"/"run_log.txt").write_text("\n".join(LOG), encoding="utf-8")
print("Done.")
